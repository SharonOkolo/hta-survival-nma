"""Build an Excel survival-extrapolation model from the R parameter estimates, and check it reproduces R.

python3 excel/build_model.py   -> excel/survival_extrapolation_model.xlsx
"""
import csv, math, os
from statistics import NormalDist
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
OUT = os.path.join(ROOT, "excel", "survival_extrapolation_model.xlsx")
MODELS = ["weibull", "lognormal", "loglogistic"]
ARMS = ["Obs", "Lev+5FU"]
MODEL_ARRAY = '{"' + '","'.join(MODELS) + '"}'  # Excel array constant for MATCH
HORIZON, STEP = 30, 1 / 12  # monthly cycles over 30 years


def load_params():
    p = {}
    with open(os.path.join(ROOT, "outputs", "parametric_params.csv")) as f:
        for r in csv.DictReader(f):
            p[(r["arm"], r["model"], r["param"])] = float(r["est"])
    return p


def surv(model, a, b, t):
    """Same formulas as the workbook. Weibull/log-logistic: a=shape, b=scale. Log-normal: a=meanlog, b=sdlog."""
    if t == 0:
        return 1.0
    if model == "weibull":
        return math.exp(-((t / b) ** a))
    if model == "loglogistic":
        return 1 / (1 + (t / b) ** a)
    return 1 - NormalDist().cdf((math.log(t) - a) / b)


def mean_survival(model, a, b):
    ts = [i * STEP for i in range(int(HORIZON / STEP) + 1)]
    s = [surv(model, a, b, t) for t in ts]
    return sum(STEP * (s[i] + s[i + 1]) / 2 for i in range(len(s) - 1))


def build(p):
    wb = Workbook()
    inp = wb.active; inp.title = "Inputs"
    bold, blue = Font(bold=True), PatternFill("solid", fgColor="DBEAFE")
    inp["A1"] = "Survival extrapolation model: colon cancer trial (Lev+5FU vs Obs)"; inp["A1"].font = Font(bold=True, size=13)
    inp["A3"], inp["B3"] = "Selected model (change me)", "lognormal"; inp["B3"].fill = blue; inp["A3"].font = bold
    dv = DataValidation(type="list", formula1='"' + ",".join(MODELS) + '"'); inp.add_data_validation(dv); dv.add("B3")
    inp["A4"], inp["B4"] = "Time horizon (years)", HORIZON
    inp.append([]); inp.append(["Model", "Arm", "Parameter 1", "Value", "Parameter 2", "Value"])
    for c in inp[6]: c.font = bold
    names = {"weibull": ("shape", "scale"), "loglogistic": ("shape", "scale"), "lognormal": ("meanlog", "sdlog")}
    rows = {}
    for m in MODELS:
        for arm in ARMS:
            n1, n2 = names[m]
            inp.append([m, arm, n1, p[(arm, m, n1)], n2, p[(arm, m, n2)]])
            rows[(m, arm)] = inp.max_row
    inp["A16"] = "Parameters estimated in R (flexsurv) by analysis/01_survival.R; source: outputs/parametric_params.csv"
    for col, w in zip("ABCDEF", (30, 12, 12, 12, 12, 12)):
        inp.column_dimensions[col].width = w

    sv = wb.create_sheet("Survival")
    sv.append(["Cycle", "Years", "S(t) Obs", "S(t) Lev+5FU", "LY Obs", "LY Lev+5FU"])
    for c in sv[1]: c.font = bold

    def par(arm, k):  # k = 1 or 2 -> pick the selected model's parameter for this arm
        col = "D" if k == 1 else "F"
        refs = ",".join(f"Inputs!${col}${rows[(m, arm)]}" for m in MODELS)
        return f"CHOOSE(MATCH(Inputs!$B$3,{MODEL_ARRAY},0),{refs})"

    def s_formula(arm, tcell):
        a, b = par(arm, 1), par(arm, 2)
        return (f'=IF({tcell}=0,1,IF(Inputs!$B$3="weibull",EXP(-({tcell}/{b})^{a}),'
                f'IF(Inputs!$B$3="loglogistic",1/(1+({tcell}/{b})^{a}),1-NORM.S.DIST((LN({tcell})-{a})/{b},TRUE))))')

    n = int(HORIZON / STEP)
    for i in range(n + 1):
        r = i + 2
        sv.append([i, f"=A{r}/12", s_formula("Obs", f"B{r}"), s_formula("Lev+5FU", f"B{r}"),
                   None if i == 0 else f"=(C{r-1}+C{r})/2/12", None if i == 0 else f"=(D{r-1}+D{r})/2/12"])
    last = n + 2
    res = wb.create_sheet("Results", 1)
    res["A1"] = "Results (update with the model selector on Inputs)"; res["A1"].font = Font(bold=True, size=13)
    res.append([]); res.append(["", "Obs", "Lev+5FU", "Difference"])
    for c in res[3]: c.font = bold
    res.append(["Mean life-years (30 y)", f"=SUM(Survival!E3:E{last})", f"=SUM(Survival!F3:F{last})", "=C4-B4"])
    res.append(["Survival at 5 years", "=Survival!C62", "=Survival!D62", "=C5-B5"])
    res.append(["Survival at 20 years", "=Survival!C242", "=Survival!D242", "=C6-B6"])
    res.column_dimensions["A"].width = 28
    ch = LineChart(); ch.title = "Extrapolated overall survival"; ch.y_axis.title = "S(t)"; ch.x_axis.title = "Cycle (months)"
    ch.add_data(Reference(sv, min_col=3, max_col=4, min_row=1, max_row=last), titles_from_data=True)
    ch.height, ch.width = 9, 18
    res.add_chart(ch, "A9")
    wb.save(OUT)


def check(p):
    """The workbook's formulas (re-implemented above) must reproduce R's mean survival within 1%."""
    r_ms = {}
    with open(os.path.join(ROOT, "outputs", "mean_survival.csv")) as f:
        for r in csv.DictReader(f):
            r_ms[r["model"]] = (float(r["obs"]), float(r["lev5fu"]))
    names = {"weibull": ("shape", "scale"), "loglogistic": ("shape", "scale"), "lognormal": ("meanlog", "sdlog")}
    out = []
    for m in MODELS:
        xl = [mean_survival(m, p[(arm, m, names[m][0])], p[(arm, m, names[m][1])]) for arm in ARMS]
        for got, want in zip(xl, r_ms[m]):
            assert abs(got - want) / want < 0.01, (m, got, want)
        out.append((m, *xl, xl[1] - xl[0]))
    return out


if __name__ == "__main__":
    p = load_params()
    build(p)
    for m, o, l, g in check(p):
        print(f"{m:12} Obs {o:.2f}  Lev+5FU {l:.2f}  gained {g:.2f} life-years (matches R within 1%)")
    print("wrote", OUT)
